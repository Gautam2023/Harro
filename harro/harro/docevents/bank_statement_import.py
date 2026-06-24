import frappe
import csv
import io
import os
from datetime import datetime


# ---------------------------------------------------------------------------
# Date helpers
# ---------------------------------------------------------------------------

def _parse_icici_date(date_str):
    """
    Parse ICICI date formats into YYYY-MM-DD string.
    Handles:
        01/03/26 11:54   -> 2026-03-01
        01/03/2026       -> 2026-03-01
        01/03/26         -> 2026-03-01
    """
    date_str = str(date_str).strip()
    if not date_str:
        raise ValueError("Date is empty")

    # Strip time portion if present
    if " " in date_str:
        date_str = date_str.split(" ")[0]

    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y"):
        try:
            return datetime.strptime(date_str, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue

    raise ValueError(f"Unrecognised date format: '{date_str}'")


# ---------------------------------------------------------------------------
# File reading — supports XLS, XLSX, and CSV
# ---------------------------------------------------------------------------

def _get_file_path(file_url):
    """Resolve a Frappe file URL to an absolute filesystem path."""
    site_path = frappe.get_site_path()
    if file_url.startswith("/private"):
        return os.path.join(site_path, "private", "files", os.path.basename(file_url))
    else:
        return os.path.join(site_path, "public", "files", os.path.basename(file_url))


def _load_rows(file_url):
    """
    Detect file type from extension and return all rows as list-of-lists.
    Supports .xls, .xlsx, .xlsm, .csv.
    """
    path = _get_file_path(file_url)
    ext  = os.path.splitext(file_url.lower())[1]

    if not os.path.exists(path):
        frappe.throw(
            f"Attached file not found on disk: {path}. "
            "Please re-attach the file and try again."
        )

    if ext == ".xls":
        return _read_rows_from_xls(path)
    elif ext in (".xlsx", ".xlsm"):
        return _read_rows_from_xlsx(path)
    else:
        return _read_rows_from_csv(path)


def _read_rows_from_xls(file_path):
    import xlrd
    wb = xlrd.open_workbook(file_path)
    ws = wb.sheet_by_index(0)
    rows = []
    for i in range(ws.nrows):
        row = []
        for j in range(ws.ncols):
            cell = ws.cell(i, j)
            if cell.ctype == 3:  # date serial
                t = xlrd.xldate_as_tuple(cell.value, wb.datemode)
                row.append(datetime(*t[:6]).strftime("%d/%m/%Y"))
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def _read_rows_from_xlsx(file_path):
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([c if c is not None else "" for c in row])
    return rows


def _read_rows_from_csv(file_path):
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            with open(file_path, newline="", encoding=enc) as f:
                return list(csv.reader(f))
        except (UnicodeDecodeError, csv.Error):
            continue
    raise ValueError(f"Cannot decode file: {file_path}")


# ---------------------------------------------------------------------------
# Core conversion logic
# ---------------------------------------------------------------------------

REQUIRED_COLS = [
    "Transaction Date",
    "Transaction Description",
    "Transaction Amount",
    "Debit / Credit",
    "Reference No.",
]

def _find_header_row(rows):
    """Scan rows until we find one whose first cell is 'Transaction Date'."""
    for i, row in enumerate(rows):
        cells = [str(c).strip() for c in row]
        if cells and cells[0] == "Transaction Date":
            col_map = {c: idx for idx, c in enumerate(cells)}
            return i, col_map
    raise ValueError(
        "Could not find the transaction header row "
        "(expected a row starting with 'Transaction Date'). "
        "Please make sure you are uploading a valid ICICI bank statement."
    )


def _convert_rows(rows, bank_account, currency):
    """
    Convert raw ICICI rows to ERPNext format.

    Returns:
        output_rows  : list of converted data rows (no header)
        skipped_rows : list of dicts — each skipped row with reason and raw data
    """
    header_idx, col_map = _find_header_row(rows)

    missing = [c for c in REQUIRED_COLS if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    idx_date   = col_map["Transaction Date"]
    idx_desc   = col_map["Transaction Description"]
    idx_amount = col_map["Transaction Amount"]
    idx_dc     = col_map["Debit / Credit"]
    idx_ref    = col_map["Reference No."]
    max_idx    = max(idx_date, idx_desc, idx_amount, idx_dc, idx_ref)

    output_rows  = []
    skipped_rows = []   # {"row": N, "reason": "...", "raw": "..."}

    for row_num, raw_row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        # Skip completely blank rows silently
        if not any(str(c).strip() for c in raw_row):
            continue

        # Row too short
        if len(raw_row) <= max_idx:
            skipped_rows.append({
                "row": row_num,
                "reason": f"Row has only {len(raw_row)} columns, expected at least {max_idx + 1}",
                "raw": ", ".join(str(c) for c in raw_row[:6]),
            })
            continue

        # --- Parse each field, collecting specific errors ---
        errors = []

        # Date
        try:
            txn_date = _parse_icici_date(str(raw_row[idx_date]))
        except Exception as e:
            txn_date = None
            errors.append(f"Date: {e}")

        # Amount
        amount_str = str(raw_row[idx_amount]).strip().replace(",", "")
        try:
            amount = float(amount_str) if amount_str else 0.0
            if amount <= 0:
                errors.append(f"Amount is zero or negative ({amount_str!r})")
        except ValueError:
            amount = None
            errors.append(f"Amount is not a valid number ({amount_str!r})")

        # D/C flag
        dc_flag = str(raw_row[idx_dc]).strip().upper()
        if dc_flag not in ("D", "C"):
            errors.append(f"Debit/Credit flag is '{dc_flag}' (expected D or C)")

        description = str(raw_row[idx_desc]).strip()
        reference   = str(raw_row[idx_ref]).strip()

        if errors:
            skipped_rows.append({
                "row": row_num,
                "reason": "; ".join(errors),
                "raw": f"{raw_row[idx_date]} | {description[:40]} | {amount_str} | {dc_flag}",
            })
            continue

        deposit    = amount if dc_flag == "C" else 0.0
        withdrawal = amount if dc_flag == "D" else 0.0

        output_rows.append(
            [txn_date, deposit, withdrawal, description, reference, bank_account, currency]
        )

    return output_rows, skipped_rows


# ---------------------------------------------------------------------------
# Frappe whitelisted entry point
# ---------------------------------------------------------------------------

@frappe.whitelist()
def process_icici_bank_statement(file_url, bank_account, currency, docname):
    """
    Convert an ICICI bank statement (XLS, XLSX, or CSV) to ERPNext import format.
    Returns detailed success/skip/error information for UI display.
    """
    try:
        rows = _load_rows(file_url)
    except frappe.exceptions.ValidationError:
        raise
    except Exception as e:
        return {"status": "error", "error": str(e)}

    try:
        output_rows, skipped_rows = _convert_rows(rows, bank_account, currency)
    except Exception as e:
        return {"status": "error", "error": str(e)}

    if not output_rows:
        return {
            "status": "error",
            "error": (
                f"No valid transaction rows found. "
                f"{len(skipped_rows)} rows were skipped. "
                "Please verify the file is a valid ICICI bank statement."
            ),
            "skipped_rows": skipped_rows,
        }

    # --- Write converted CSV ---
    output_buffer = io.StringIO()
    writer = csv.writer(output_buffer)
    writer.writerow(["Date", "Deposit", "Withdrawal", "Description",
                     "Reference Number", "Bank Account", "Currency"])
    writer.writerows(output_rows)
    csv_content = output_buffer.getvalue()

    safe_name       = docname.replace(" ", "_").replace(":", "").replace("/", "")
    output_filename = f"{safe_name}_converted.csv"

    try:
        file_doc = frappe.get_doc({
            "doctype": "File",
            "file_name": output_filename,
            "attached_to_doctype": "Bank Statement Import",
            "attached_to_name": docname,
            "attached_to_field": "import_file",
            "content": csv_content.encode("utf-8"),
            "is_private": 1,
        })
        file_doc.insert(ignore_permissions=True)
    except Exception as e:
        return {"status": "error", "error": f"Failed to save converted file: {e}"}

    try:
        frappe.db.set_value("Bank Statement Import", docname,
                            "import_file", file_doc.file_url)
        frappe.db.commit()
    except Exception as e:
        return {
            "status": "error",
            "error": f"File saved ({file_doc.file_url}) but could not link to import_file: {e}",
        }

    # Log skipped rows to Error Log for developer visibility
    if skipped_rows:
        frappe.log_error(
            "\n".join(
                f"Row {r['row']}: {r['reason']} | Raw: {r['raw']}"
                for r in skipped_rows
            ),
            f"ICICI Statement Skipped Rows — {docname}",
        )

    return {
        "status": "success",
        "rows": len(output_rows),
        "skipped": len(skipped_rows),
        "skipped_rows": skipped_rows,   # full detail for UI
        "file_url": file_doc.file_url,
    }